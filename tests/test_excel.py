"""Checks Excel import and export. Uses a temporary database, never your real data.
Run from the project folder: python tests/test_excel.py
"""
import io
import sys
import tempfile
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import spendwise.database as db

db.DB_PATH = Path(tempfile.mkdtemp()) / "test.db"
db.init_db()

from openpyxl import Workbook, load_workbook  # noqa: E402

from spendwise.services import csv_io, transactions as tx  # noqa: E402

failures = 0


def check(label, actual, expected):
    global failures
    ok = actual == expected
    print("PASS" if ok else "FAIL", "-", label, "" if ok else f"(expected {expected!r}, got {actual!r})")
    if not ok:
        failures += 1


def make_workbook(rows, headers=("Date", "Type", "Category", "Amount", "Description")):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


ids = {(c["name"], c["type"]): c["id"] for c in tx.get_categories()}
FOOD = ids[("Food & Dining", "expense")]
TRANSPORT = ids[("Transport", "expense")]
SALARY = ids[("Salary", "income")]

# One saved row, used to check duplicate detection.
tx.add_transaction("2026-09-30", "expense", 99, FOOD, "Existing")

rows = [
    [datetime(2026, 10, 1), "income", "Salary", 50000, "October salary"],
    ["23/10/2026", "Expense", "Food & Dining", 250.5, "Lunch"],
    [None, None, None, None, None],
    ["not-a-date", "expense", "Food & Dining", 10, ""],
    ["2026-10-02", "expense", "Nope", 5, ""],
    ["2026-09-30", "expense", "Food & Dining", 99, "Existing"],
    ["2026-10-04", "expense", "Transport", "1,250.50", ""],
]
data = make_workbook(rows)

print("\nReading an Excel file")
result = csv_io.parse_xlsx(data)
check("rows read (blank row skipped)", result["total_rows"], 6)
check("new rows", len(result["new"]), 3)
check("duplicate rows", len(result["duplicates"]), 1)
check("problem rows", len(result["errors"]), 2)
check("problem row numbers", [e["row"] for e in result["errors"]], [5, 6])
check(
    "problem messages",
    [e["message"] for e in result["errors"]],
    [
        "Date 'not-a-date' is not valid. Use YYYY-MM-DD or DD/MM/YYYY.",
        "Unknown expense category 'Nope'.",
    ],
)
check("real Excel date is converted", result["new"][0]["date"], "2026-10-01")
check("text date still works", result["new"][1]["date"], "2026-10-23")
check("numeric amount", result["new"][1]["amount"], 250.5)
check("amount with a comma", result["new"][2]["amount"], 1250.5)
check("type is normalised", result["new"][1]["type"], "expense")
check("category is normalised", result["new"][1]["category"], "Food & Dining")

print("\nChoosing the parser")
check("xlsx by extension", len(csv_io.parse_file(data, "Data.XLSX")["new"]), 3)
csv_data = b"date,type,category,amount,description\n2026-10-05,expense,Transport,20,Train\n"
check("csv by extension", len(csv_io.parse_file(csv_data, "data.csv")["new"]), 1)

print("\nBad files")
try:
    csv_io.parse_xlsx(make_workbook([[1, 2, 3]], headers=("a", "b", "c")))
    check("missing columns are rejected", "accepted", "rejected")
except tx.ValidationError:
    check("missing columns are rejected", "rejected", "rejected")
try:
    csv_io.parse_xlsx(b"this is not a spreadsheet")
    check("broken file is rejected", "accepted", "rejected")
except tx.ValidationError:
    check("broken file is rejected", "rejected", "rejected")

print("\nExporting")
exported = csv_io.export_xlsx()
check("returns bytes", isinstance(exported, bytes), True)
sheet = load_workbook(io.BytesIO(exported)).active
check("sheet name", sheet.title, "Transactions")
check("header row", [c.value for c in sheet[1]], ["date", "type", "category", "amount", "description"])
check("header is bold", sheet["A1"].font.bold, True)
check("header is frozen", sheet.freeze_panes, "A2")
check("amount is a real number", isinstance(sheet["D2"].value, (int, float)), True)
check("amount is formatted", sheet["D2"].number_format, "#,##0.00")

roundtrip = csv_io.parse_xlsx(exported)
check(
    "an export re-imports as duplicates",
    (len(roundtrip["new"]), len(roundtrip["duplicates"]), len(roundtrip["errors"])),
    (0, 1, 0),
)

print("\nFormula-looking descriptions stay text")
tx.add_transaction("2026-10-06", "expense", 12, FOOD, "=SUM(A1:A2)")
exported = csv_io.export_xlsx()
sheet = load_workbook(io.BytesIO(exported)).active
formula_cell = next(c for row in sheet.iter_rows(min_row=2) for c in row if c.column == 5 and c.value == "=SUM(A1:A2)")
check("stored as text, not a formula", formula_cell.data_type, "s")
reimported = csv_io.parse_xlsx(exported)
check(
    "round-trip keeps the description",
    "=SUM(A1:A2)" in [r["description"] for r in reimported["duplicates"]],
    True,
)

print("\nTemplate")
template = csv_io.template_xlsx()
check("template is a workbook", load_workbook(io.BytesIO(template)).active["A1"].value, "date")
parsed_template = csv_io.parse_xlsx(template)
check("template has two usable rows", len(parsed_template["new"]), 2)
check("template has no problems", parsed_template["errors"], [])

print()
print("All checks passed." if failures == 0 else f"{failures} check(s) FAILED.")
sys.exit(1 if failures else 0)
