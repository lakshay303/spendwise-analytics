import csv
import io
import re
import zipfile
from collections import Counter
from datetime import date, datetime

from spendwise.database import get_connection, query
from spendwise.services.transactions import VALID_TYPES, ValidationError, _clean, list_transactions

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
MAX_UNZIPPED_BYTES = 50 * 1024 * 1024
EXCEL_ROW_SLACK = 1000  # allowance for blank rows in a sheet
COLUMNS = ["date", "type", "category", "amount", "description"]
REQUIRED = ["date", "type", "category", "amount"]
DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y")
TEMPLATE_ROWS = [
    ["2026-10-01", "income", "Salary", 50000.00, "October salary"],
    ["2026-10-02", "expense", "Food & Dining", 250.00, "Lunch"],
]


def _protect(text):
    """Stop spreadsheet programs from running a description that starts like a formula."""
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def _unprotect(text):
    """Undo _protect when re-importing an exported file."""
    if len(text) > 1 and text[0] == "'" and text[1] in "=+-@":
        return text[1:]
    return text


def _show(value, limit=40):
    """Shorten user text before putting it in a message."""
    value = str(value)
    return value if len(value) <= limit else value[:limit] + "…"


def _parse_date(text):
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _decode(data):
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValidationError("Could not read the file. Save it as CSV (UTF-8) and try again.")


def _cell(value):
    """Turn one spreadsheet cell into a clean string. Empty cells become ''."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    try:
        if value != value:  # NaN or NaT
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _missing_columns(names):
    missing = [name for name in REQUIRED if name not in names]
    if missing:
        raise ValidationError(
            "The file is missing required column(s): " + ", ".join(missing)
            + ". Expected columns: date, type, category, amount, and optionally description."
        )


def _classify(records):
    """Sort (row_number, values) records into new, duplicates and errors. Saves nothing.

    A record with values=None means the row could not be read (CSV row with too many columns).
    """
    categories = {
        (c["name"].lower(), c["type"]): (c["id"], c["name"])
        for c in query("SELECT id, name, type FROM categories")
    }
    existing = Counter(
        (r["date"], r["type"], round(r["amount"], 2), r["category_id"], r["description"])
        for r in list_transactions()
    )

    new, duplicates, errors = [], [], []
    total_rows = 0
    for line_no, row in records:
        if row is not None and all(not value for value in row.values()):
            continue  # completely blank row

        total_rows += 1
        if total_rows > MAX_ROWS:
            raise ValidationError(f"Too many rows. The limit is {MAX_ROWS} per file.")
        if row is None:
            errors.append({"row": line_no, "message": "This row has more columns than the header."})
            continue

        values, problem = _check_row(row, categories)
        if problem:
            errors.append({"row": line_no, "message": problem})
            continue

        values["row"] = line_no
        key = (values["date"], values["type"], values["amount"], values["category_id"], values["description"])
        if existing[key] > 0:
            existing[key] -= 1
            duplicates.append(values)
        else:
            new.append(values)

    return {"new": new, "duplicates": duplicates, "errors": errors, "total_rows": total_rows}


def export_csv(rows=None):
    """Return transactions as CSV text. Pass rows (from list_transactions) to export a filtered set."""
    if rows is None:
        rows = list_transactions(sort_by="date", descending=False)
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(COLUMNS)
    for r in rows:
        writer.writerow([r["date"], r["type"], r["category"], f"{r['amount']:.2f}", _protect(r["description"])])
    return buffer.getvalue()


def _write_workbook(rows):
    """Build an .xlsx file (bytes) with a styled header. rows: [date, type, category, amount, description]."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transactions"

    sheet.append(COLUMNS)
    for row in rows:
        sheet.append(row)
        # Write descriptions as text, never as a formula (openpyxl treats a leading "=" as one).
        description_cell = sheet.cell(row=sheet.max_row, column=COLUMNS.index("description") + 1)
        description_cell.data_type = "s"

    header_fill = PatternFill("solid", fgColor="4F46E5")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    amount_column = COLUMNS.index("amount") + 1
    for (cell,) in sheet.iter_rows(min_row=2, min_col=amount_column, max_col=amount_column):
        cell.number_format = "#,##0.00"

    widths = {"date": 12, "type": 10, "category": 20, "amount": 14, "description": 40}
    for name, width in widths.items():
        sheet.column_dimensions[get_column_letter(COLUMNS.index(name) + 1)].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_xlsx(rows=None):
    """Return transactions as Excel (.xlsx) bytes. Pass rows to export a filtered set."""
    if rows is None:
        rows = list_transactions(sort_by="date", descending=False)
    return _write_workbook(
        [[r["date"], r["type"], r["category"], float(r["amount"]), r["description"]] for r in rows]
    )


def template_xlsx():
    """Return an Excel file (bytes) with the column headers and two sample rows."""
    return _write_workbook(TEMPLATE_ROWS)


def parse_csv(data):
    """Read CSV bytes and sort the rows into new, duplicates and errors. Saves nothing."""
    if len(data) > MAX_BYTES:
        raise ValidationError("The file is too large. The limit is 5 MB.")
    text = _decode(data).lstrip()
    if not text:
        raise ValidationError("The file is empty.")

    first_line = text.splitlines()[0]
    delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    reader.fieldnames = [(name or "").strip().lower() for name in (reader.fieldnames or [])]
    _missing_columns(reader.fieldnames)

    def records():
        for line_no, row in enumerate(reader, start=2):
            yield line_no, (None if None in row else {name: _cell(value) for name, value in row.items()})

    return _classify(records())


def _check_unpacked_size(data):
    """Refuse .xlsx files that are small to upload but huge once unpacked."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            unpacked = sum(item.file_size for item in archive.infolist())
    except zipfile.BadZipFile:
        raise ValidationError(
            "Could not read the Excel file. Save it as .xlsx with the column names in the first row."
        )
    if unpacked > MAX_UNZIPPED_BYTES:
        raise ValidationError("The Excel file is too large once unpacked. Remove unused data and try again.")


def parse_xlsx(data):
    """Read the first sheet of an .xlsx file and sort the rows into new, duplicates and errors."""
    if len(data) > MAX_BYTES:
        raise ValidationError("The file is too large. The limit is 5 MB.")
    if not data:
        raise ValidationError("The file is empty.")
    _check_unpacked_size(data)

    import pandas as pd

    try:
        frame = pd.read_excel(
            io.BytesIO(data),
            sheet_name=0,
            dtype=object,
            nrows=MAX_ROWS + EXCEL_ROW_SLACK + 1,
        )
    except Exception:
        raise ValidationError(
            "Could not read the Excel file. Save it as .xlsx with the column names in the first row."
        )
    if len(frame) > MAX_ROWS + EXCEL_ROW_SLACK:
        raise ValidationError(f"Too many rows. The limit is {MAX_ROWS} per file.")

    column_map = {}
    for name in frame.columns:
        column_map.setdefault(str(name).strip().lower(), name)
    _missing_columns(column_map)

    def records():
        for position, row in enumerate(frame.itertuples(index=False), start=2):
            values = dict(zip(frame.columns, row))
            yield position, {name: _cell(values[column_map[name]]) for name in COLUMNS if name in column_map}

    return _classify(records())

def parse_file(data, filename):
    """Read an uploaded .csv or .xlsx file. Returns the same result as parse_csv()."""
    if (filename or "").lower().endswith((".xlsx", ".xlsm")):
        return parse_xlsx(data)
    return parse_csv(data)


def _check_row(row, categories):
    """Return (values, None) for a good row, or (None, message) for a bad one."""
    raw_date = (row.get("date") or "").strip()
    tx_type = (row.get("type") or "").strip().lower()
    category = (row.get("category") or "").strip()
    raw_amount = (row.get("amount") or "").strip()
    description = _unprotect((row.get("description") or "").strip())

    if not raw_date:
        return None, "Date is missing."
    iso_date = _parse_date(raw_date)
    if not iso_date:
        return None, f"Date '{_show(raw_date)}' is not valid. Use YYYY-MM-DD or DD/MM/YYYY."
    if tx_type not in VALID_TYPES:
        return None, f"Type '{_show(tx_type)}' must be 'income' or 'expense'."
    if not category:
        return None, "Category is missing."
    if not raw_amount:
        return None, "Amount is missing."
    try:
        amount = float(re.sub(r"[₹,\s]", "", raw_amount))
    except ValueError:
        return None, f"Amount '{_show(raw_amount)}' is not a number."

    match = categories.get((category.lower(), tx_type))
    if not match:
        return None, f"Unknown {tx_type} category '{_show(category)}'."
    category_id, category_name = match

    try:
        iso_date, tx_type, amount, category_id, description = _clean(
            iso_date, tx_type, amount, category_id, description
        )
    except ValidationError as error:
        return None, str(error)

    return {
        "date": iso_date,
        "type": tx_type,
        "category": category_name,
        "category_id": category_id,
        "amount": amount,
        "description": description,
    }, None


def import_rows(rows):
    """Save rows from parse_csv()['new'] in one go: either all are saved or none. Returns the count."""
    if not rows:
        return 0
    with get_connection() as conn:
        conn.executemany(
            "INSERT INTO transactions (date, type, amount, category_id, description) VALUES (?, ?, ?, ?, ?)",
            [(r["date"], r["type"], r["amount"], r["category_id"], r["description"]) for r in rows],
        )
    return len(rows)
